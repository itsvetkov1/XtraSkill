"""
Document management routes.

Handles file upload, listing, viewing, and search for project documents.
"""

import csv
import json
from io import BytesIO, StringIO
from typing import List, Optional
from fastapi import APIRouter, Depends, File, HTTPException, Response, status, UploadFile
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload
import openpyxl

from app.database import get_db
from app.models import Document, Project, Thread, User
from app.routes.auth import get_current_user
from app.services.encryption import get_encryption_service
from app.services.document_search import index_document, search_documents
from app.services.document_parser import ParserFactory
from app.services.file_validator import validate_file_security


router = APIRouter()

# File upload constraints — rich documents up to 10MB
ALLOWED_CONTENT_TYPES = ParserFactory.ALL_CONTENT_TYPES

# Tabular formats that support export
TABULAR_CONTENT_TYPES = [
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "text/csv",
]


async def _process_and_store_document(
    db: AsyncSession,
    file: UploadFile,
    project_id: Optional[str],
    thread_id: Optional[str]
) -> Document:
    """
    Process, encrypt, and store a document.

    Shared logic for both project and thread document uploads.

    Args:
        db: Database session
        file: Uploaded file
        project_id: Optional project ID (for project documents)
        thread_id: Optional thread ID (for thread documents)

    Returns:
        Created Document record

    Raises:
        HTTPException: For validation or processing errors
    """
    # Validate file type
    if file.content_type not in ALLOWED_CONTENT_TYPES:
        raise HTTPException(
            status_code=400,
            detail="Unsupported file type. Supported: .txt, .md, .xlsx, .csv, .pdf, .docx"
        )

    # Read file bytes
    content_bytes = await file.read()

    # Security validation (size, magic numbers, zip bombs)
    validate_file_security(content_bytes, file.content_type)

    # Get parser for content type
    parser = ParserFactory.get_parser(file.content_type)

    # Format-specific security validation
    parser.validate_security(content_bytes)

    # Parse document to extract text and metadata
    try:
        parsed = parser.parse(content_bytes)
    except Exception as e:
        raise HTTPException(
            status_code=400,
            detail=f"Failed to parse document: {str(e)}"
        )

    # Encrypt for storage
    if ParserFactory.is_rich_format(file.content_type):
        # Rich formats — store original binary
        encrypted = get_encryption_service().encrypt_binary(content_bytes)
    else:
        # Text formats — store plaintext (backward compatible)
        encrypted = get_encryption_service().encrypt_document(parsed["text"])

    # Create document record with dual-column storage
    doc = Document(
        project_id=project_id,
        thread_id=thread_id,
        filename=file.filename or "untitled",
        content_type=file.content_type,
        content_encrypted=encrypted,
        content_text=parsed["text"],
        metadata_json=json.dumps(parsed["metadata"]) if parsed["metadata"] else None,
    )
    db.add(doc)
    await db.flush()  # Get doc.id

    # Index for search (in same transaction)
    await index_document(db, doc.id, doc.filename, parsed["text"])

    return doc


@router.post("/projects/{project_id}/documents", status_code=201)
async def upload_document(
    project_id: str,
    file: UploadFile = File(...),
    current_user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db)
):
    """
    Upload a document to a project.

    Supports .txt, .md, .xlsx, .csv, .pdf, .docx files, max 10MB.
    Content is encrypted at rest and indexed for full-text search.
    """
    # Verify project exists and belongs to current user
    stmt = select(Project).where(
        Project.id == project_id,
        Project.user_id == current_user["user_id"]
    )
    project = (await db.execute(stmt)).scalar_one_or_none()
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")

    # Process and store document
    doc = await _process_and_store_document(db, file, project_id=project_id, thread_id=None)

    await db.commit()

    # Parse metadata for response
    metadata = json.loads(doc.metadata_json) if doc.metadata_json else None

    return {
        "id": doc.id,
        "filename": doc.filename,
        "content_type": doc.content_type,
        "metadata": metadata,
        "created_at": doc.created_at.isoformat()
    }


@router.get("/projects/{project_id}/documents")
async def list_documents(
    project_id: str,
    current_user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db)
):
    """
    List all documents in a project.

    Returns metadata only (no content) ordered by creation date.
    """
    # Verify project ownership
    stmt = select(Project).where(
        Project.id == project_id,
        Project.user_id == current_user["user_id"]
    )
    project = (await db.execute(stmt)).scalar_one_or_none()
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")

    # Get documents
    stmt = select(Document).where(
        Document.project_id == project_id
    ).order_by(Document.created_at.desc())

    result = await db.execute(stmt)
    documents = result.scalars().all()

    return [
        {
            "id": doc.id,
            "filename": doc.filename,
            "content_type": doc.content_type or "text/plain",
            "metadata": json.loads(doc.metadata_json) if doc.metadata_json else None,
            "created_at": doc.created_at.isoformat()
        }
        for doc in documents
    ]


@router.get("/documents/{document_id}")
async def get_document(
    document_id: str,
    current_user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db)
):
    """
    Get document content.

    Verifies user owns the project containing the document.
    Returns extracted text content for all document types.
    """
    # Get document with project join to verify ownership
    stmt = select(Document).join(Project).where(
        Document.id == document_id,
        Project.user_id == current_user["user_id"]
    )
    doc = (await db.execute(stmt)).scalar_one_or_none()
    if not doc:
        raise HTTPException(status_code=404, detail="Document not found")

    # Get content text
    if doc.content_text is not None:
        content = doc.content_text
    else:
        # Legacy document — decrypt from content_encrypted (text format)
        content = get_encryption_service().decrypt_document(doc.content_encrypted)

    return {
        "id": doc.id,
        "filename": doc.filename,
        "content_type": doc.content_type or "text/plain",
        "content": content,
        "metadata": json.loads(doc.metadata_json) if doc.metadata_json else None,
        "created_at": doc.created_at.isoformat()
    }


@router.get("/documents/{document_id}/download")
async def download_document(
    document_id: str,
    current_user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db)
):
    """
    Download original document file.

    For rich documents (XLSX, CSV, PDF, DOCX): returns original binary.
    For text documents: returns plain text content.
    """
    # Get document with project join to verify ownership
    stmt = select(Document).join(Project).where(
        Document.id == document_id,
        Project.user_id == current_user["user_id"]
    )
    doc = (await db.execute(stmt)).scalar_one_or_none()
    if not doc:
        raise HTTPException(status_code=404, detail="Document not found")

    content_type = doc.content_type or "text/plain"

    if ParserFactory.is_rich_format(content_type):
        # Rich document — decrypt as binary
        file_bytes = get_encryption_service().decrypt_binary(doc.content_encrypted)
        return Response(
            content=file_bytes,
            media_type=content_type,
            headers={
                "Content-Disposition": f'attachment; filename="{doc.filename}"'
            }
        )
    else:
        # Text document — decrypt as text
        plaintext = get_encryption_service().decrypt_document(doc.content_encrypted)
        return Response(
            content=plaintext.encode('utf-8'),
            media_type=content_type,
            headers={
                "Content-Disposition": f'attachment; filename="{doc.filename}"'
            }
        )


async def _get_tabular_document(
    document_id: str,
    current_user: dict,
    db: AsyncSession
) -> Document:
    """
    Helper function to get a document and verify it's tabular format.

    Args:
        document_id: Document UUID
        current_user: Authenticated user
        db: Database session

    Returns:
        Document object if found and tabular

    Raises:
        404: Document not found or not owned by user
        400: Document is not tabular format or has no data
    """
    # Get document with project join to verify ownership
    stmt = select(Document).join(Project).where(
        Document.id == document_id,
        Project.user_id == current_user["user_id"]
    )
    doc = (await db.execute(stmt)).scalar_one_or_none()
    if not doc:
        raise HTTPException(status_code=404, detail="Document not found")

    # Check if tabular format
    if doc.content_type not in TABULAR_CONTENT_TYPES:
        raise HTTPException(
            status_code=400,
            detail="Export is only available for spreadsheet documents"
        )

    # Check if content exists
    if not doc.content_text:
        raise HTTPException(
            status_code=400,
            detail="No data available for export"
        )

    return doc


@router.get("/documents/{document_id}/export/xlsx")
async def export_document_xlsx(
    document_id: str,
    current_user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db)
):
    """
    Export document as Excel (.xlsx) file.

    Only available for spreadsheet documents (Excel, CSV).
    Returns in-memory generated Excel file with proper headers for browser download.
    """
    # Get and validate document
    doc = await _get_tabular_document(document_id, current_user, db)

    # Parse metadata for sheet name
    metadata = json.loads(doc.metadata_json) if doc.metadata_json else {}
    sheet_name = metadata.get('sheet_names', ['Sheet1'])[0]

    # Generate Excel in memory
    output = BytesIO()
    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet(title=sheet_name)

    # Parse content_text (tab-separated) and write rows
    for line in doc.content_text.split('\n'):
        if line.strip():  # Skip empty lines
            cells = line.split('\t')
            ws.append(cells)

    wb.save(output)
    output.seek(0)

    # Generate filename
    base_name = doc.filename.rsplit('.', 1)[0] if '.' in doc.filename else doc.filename
    export_filename = f"{base_name}_export.xlsx"

    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="{export_filename}"'
        }
    )


@router.get("/documents/{document_id}/export/csv")
async def export_document_csv(
    document_id: str,
    current_user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db)
):
    """
    Export document as CSV file.

    Only available for spreadsheet documents (Excel, CSV).
    Returns in-memory generated CSV with UTF-8 BOM for Excel compatibility.
    """
    # Get and validate document
    doc = await _get_tabular_document(document_id, current_user, db)

    # Generate CSV in memory
    output = StringIO()
    writer = csv.writer(output, quoting=csv.QUOTE_MINIMAL)

    # Parse content_text (tab-separated) and write rows
    for line in doc.content_text.split('\n'):
        if line.strip():  # Skip empty lines
            cells = line.split('\t')
            writer.writerow(cells)

    # Encode with UTF-8 BOM for Excel compatibility
    csv_bytes = output.getvalue().encode('utf-8-sig')

    # Generate filename
    base_name = doc.filename.rsplit('.', 1)[0] if '.' in doc.filename else doc.filename
    export_filename = f"{base_name}_export.csv"

    return Response(
        content=csv_bytes,
        media_type="text/csv; charset=utf-8",
        headers={
            "Content-Disposition": f'attachment; filename="{export_filename}"'
        }
    )


@router.get("/projects/{project_id}/documents/search")
async def search_project_documents(
    project_id: str,
    q: str,
    current_user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db)
):
    """
    Search documents within a project using full-text search.

    Returns ranked results with snippets and BM25 relevance scores.
    """
    # Verify project ownership
    stmt = select(Project).where(
        Project.id == project_id,
        Project.user_id == current_user["user_id"]
    )
    project = (await db.execute(stmt)).scalar_one_or_none()
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")

    # Search documents
    results = await search_documents(db, project_id, q)

    return [
        {
            "id": doc_id,
            "filename": filename,
            "snippet": snippet,
            "score": score
        }
        for doc_id, filename, snippet, score in results
    ]


@router.post("/threads/{thread_id}/documents", status_code=201)
async def upload_thread_document(
    thread_id: str,
    file: UploadFile = File(...),
    current_user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db)
):
    """
    Upload a document to a thread (Assistant mode).

    Supports .txt, .md, .xlsx, .csv, .pdf, .docx files, max 10MB.
    Content is encrypted at rest and indexed for full-text search.
    Thread documents are not associated with projects.

    Args:
        thread_id: Thread UUID
        file: Uploaded file
        current_user: Authenticated user from JWT
        db: Database session

    Returns:
        201 Created with document metadata

    Raises:
        404: Thread not found or not owned by user
        400: Invalid file type or processing error
    """
    user_id = current_user["user_id"]

    # Get thread with project loaded for ownership check
    stmt = (
        select(Thread)
        .where(Thread.id == thread_id)
        .options(selectinload(Thread.project))
    )
    result = await db.execute(stmt)
    thread = result.scalar_one_or_none()

    if not thread:
        raise HTTPException(status_code=404, detail="Thread not found")

    # Validate ownership: project-less threads check user_id, project threads check project.user_id
    if thread.project_id is None:
        if thread.user_id != user_id:
            raise HTTPException(status_code=404, detail="Thread not found")
    else:
        if thread.project.user_id != user_id:
            raise HTTPException(status_code=404, detail="Thread not found")

    # Process and store document (no project_id for thread documents)
    doc = await _process_and_store_document(db, file, project_id=None, thread_id=thread_id)

    await db.commit()

    # Parse metadata for response
    metadata = json.loads(doc.metadata_json) if doc.metadata_json else None

    return {
        "id": doc.id,
        "filename": doc.filename,
        "content_type": doc.content_type,
        "metadata": metadata,
        "created_at": doc.created_at.isoformat()
    }


@router.get("/threads/{thread_id}/documents")
async def list_thread_documents(
    thread_id: str,
    current_user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db)
):
    """
    List all documents uploaded to a thread.

    Returns metadata only (no content) ordered by creation date (newest first).

    Args:
        thread_id: Thread UUID
        current_user: Authenticated user from JWT
        db: Database session

    Returns:
        List of document metadata

    Raises:
        404: Thread not found or not owned by user
    """
    user_id = current_user["user_id"]

    # Get thread with project loaded for ownership check
    stmt = (
        select(Thread)
        .where(Thread.id == thread_id)
        .options(selectinload(Thread.project))
    )
    result = await db.execute(stmt)
    thread = result.scalar_one_or_none()

    if not thread:
        raise HTTPException(status_code=404, detail="Thread not found")

    # Validate ownership: project-less threads check user_id, project threads check project.user_id
    if thread.project_id is None:
        if thread.user_id != user_id:
            raise HTTPException(status_code=404, detail="Thread not found")
    else:
        if thread.project.user_id != user_id:
            raise HTTPException(status_code=404, detail="Thread not found")

    # Get documents for this thread
    stmt = select(Document).where(
        Document.thread_id == thread_id
    ).order_by(Document.created_at.desc())

    result = await db.execute(stmt)
    documents = result.scalars().all()

    return [
        {
            "id": doc.id,
            "filename": doc.filename,
            "content_type": doc.content_type or "text/plain",
            "metadata": json.loads(doc.metadata_json) if doc.metadata_json else None,
            "created_at": doc.created_at.isoformat()
        }
        for doc in documents
    ]


@router.delete("/documents/{document_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_document(
    document_id: str,
    current_user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db)
):
    """
    Delete a document.

    Verifies user owns the project containing the document.

    Args:
        document_id: Document UUID
        current_user: Authenticated user from JWT
        db: Database session

    Returns:
        204 No Content on success

    Raises:
        404: Document not found or not owned by user
    """
    # Get document with project join to verify ownership
    stmt = select(Document).join(Project).where(
        Document.id == document_id,
        Project.user_id == current_user["user_id"]
    )
    doc = (await db.execute(stmt)).scalar_one_or_none()
    if not doc:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Document not found"
        )

    # Delete document
    await db.delete(doc)
    await db.commit()

    return Response(status_code=status.HTTP_204_NO_CONTENT)
